"""HTTP workflows over an isolated SQLite database; no production startup."""
import asyncio
import io
import json
import shutil
import unittest
import uuid
import zipfile
from pathlib import Path
from unittest.mock import patch

import httpx
from PIL import Image
from openpyxl import load_workbook
from sqlmodel import SQLModel, Session, create_engine

from app.main import app
from app.api.deps import current_active_user, get_async_session
from app.core.config import settings
from app.models.analysis import Analysis
from app.models.user import UserTable
from app.features.analysis import queue
from app.features.analysis import service as analysis_service
from app.features.export import excel
from app.features.export.responses import XLSX_MEDIA_TYPE
from app.infrastructure.storage.previews import ensure_preview, preview_path_for


class AsyncSessionAdapter:
    """Run the real SQLAlchemy statements on the test's synchronous SQLite DB."""
    def __init__(self, session):
        self.session = session

    async def execute(self, statement):
        return self.session.execute(statement)

    def add(self, record):
        self.session.add(record)

    async def commit(self):
        self.session.commit()

    async def refresh(self, record):
        self.session.refresh(record)

    async def delete(self, record):
        self.session.delete(record)


class HttpWorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        run_dir = Path(__file__).resolve().parents[2] / ".run"
        self.root = run_dir / ("http-test-" + uuid.uuid4().hex)
        self.root.mkdir()
        self.engine = create_engine("sqlite:///" + str(self.root / "test.db"))
        SQLModel.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        self.owner = UserTable(email="owner@example.com", hashed_password="test-only")
        self.other = UserTable(email="other@example.com", hashed_password="test-only")
        self.session.add_all([self.owner, self.other])
        self.session.commit()
        self.session.refresh(self.owner)
        self.session.refresh(self.other)
        self.owner_record = self.make_record(self.owner, "sample", "stem")
        self.other_record = self.make_record(self.other, "foreign", "leaf")
        self.overrides = app.dependency_overrides.copy()

        async def test_session():
            yield AsyncSessionAdapter(self.session)

        async def test_user():
            return self.owner

        app.dependency_overrides[get_async_session] = test_session
        app.dependency_overrides[current_active_user] = test_user
        self.patches = [
            patch.object(settings, "STORAGE_PATH", self.root.as_posix()),
            patch.object(excel, "sync_engine", self.engine),
            patch.object(queue, "enqueue_analysis", return_value=1),
            patch.object(analysis_service, "sync_engine", self.engine),
        ]
        static_app = next(route.app for route in app.routes if getattr(route, "path", None) == "/static")
        self.patches.extend([
            patch.object(static_app, "directory", str(self.root)),
            patch.object(static_app, "all_directories", [str(self.root)]),
        ])
        for p in self.patches:
            p.start()
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")

    async def asyncTearDown(self):
        await self.client.aclose()
        app.dependency_overrides.clear()
        app.dependency_overrides.update(self.overrides)
        for p in reversed(self.patches):
            p.stop()
        self.session.close()
        self.engine.dispose()
        self.assertEqual(self.root.parent.name, ".run")
        self.assertTrue(self.root.name.startswith("http-test-"))
        shutil.rmtree(self.root)

    def make_record(self, owner, name, task_type):
        original = self.root / (name + ".png")
        Image.new("RGB", (32, 32), "white").save(original)
        ensure_preview(str(original))
        result = self.root / (name + ".json")
        result.write_text(json.dumps({
            "imageWidth": 32, "imageHeight": 32,
            "shapes": [{"label": "out", "points": [[0, 0], [4, 0], [4, 4], [0, 4]]}],
        }), encoding="utf-8")
        preview = self.root / (name + "-annotated.jpg")
        Image.new("RGB", (32, 32), "white").save(preview)
        record = Analysis(user_id=owner.id, task_type=task_type, status="completed",
                          original_file_path=str(original), result_json_path=str(result),
                          annotated_image_path=str(preview))
        self.session.add(record)
        self.session.commit()
        self.session.refresh(record)
        return record

    async def test_register_login_and_authenticated_history(self):
        app.dependency_overrides.pop(current_active_user)
        unauthenticated = await self.client.get("/api/analysis/history")
        self.assertEqual(unauthenticated.status_code, 401)
        response = await self.client.post("/api/auth/register", json={
            "email": "isolated@example.com", "password": "OnlyForThisTest123!",
        })
        self.assertEqual(response.status_code, 201, response.text)
        login = await self.client.post("/api/auth/jwt/login", data={
            "username": "isolated@example.com", "password": "OnlyForThisTest123!",
        })
        self.assertEqual(login.status_code, 200, login.text)
        history = await self.client.get("/api/analysis/history", headers={
            "Authorization": "Bearer " + login.json()["access_token"],
        })
        self.assertEqual(history.status_code, 200, history.text)
        self.assertEqual(history.json(), [])

    async def test_history_and_stats_do_not_expose_another_users_records(self):
        response = await self.client.get("/api/analysis/history")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual([r["analysisId"] for r in response.json()], [str(self.owner_record.analysis_id)])
        self.assertIn("/static/sample.preview.jpg", response.json()[0]["originalImageUrl"])
        image = await self.client.get(response.json()[0]["originalImageUrl"])
        self.assertEqual(image.status_code, 200)
        self.assertEqual(image.headers["content-type"], "image/jpeg")
        filtered = await self.client.get("/api/analysis/history?group=leaf")
        self.assertEqual(filtered.json(), [])
        stats = await self.client.get("/api/analysis/stats?days=7")
        self.assertEqual(stats.status_code, 200, stats.text)
        self.assertEqual(stats.json()["total"], 1)
        self.assertEqual(stats.json()["windowDays"], 7)

    async def test_single_and_batch_upload_persist_and_enqueue(self):
        buffer = io.BytesIO()
        Image.new("RGB", (16, 16), "white").save(buffer, format="PNG")
        png = buffer.getvalue()
        response = await self.client.post("/api/analysis/upload", files={
            "file": ("upload.png", png, "image/png"),
        }, data={"task_type": "leaf_our"})
        self.assertEqual(response.status_code, 200, response.text)
        record = self.session.get(Analysis, uuid.UUID(response.json()["analysis_id"]))
        self.assertEqual(record.status, "queued")
        self.assertEqual(record.task_type, "leaf_our")
        self.assertTrue(Path(preview_path_for(record.original_file_path)).exists())
        batch = await self.client.post("/api/analysis/upload/batch", files=[
            ("files", ("one.png", png, "image/png")),
            ("files", ("two.png", png, "image/png")),
        ], data={"task_type": "stem"})
        self.assertEqual(batch.status_code, 200, batch.text)
        self.assertEqual(batch.json()["total_files"], 2)
        self.assertEqual(queue.enqueue_analysis.call_count, 3)

    async def test_json_and_excel_exports_preserve_content(self):
        record_id = str(self.owner_record.analysis_id)
        response = await self.client.get("/api/export/json/" + record_id)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "application/octet-stream")
        self.assertEqual(response.json()["shapes"][0]["label"], "out")
        zipped = await self.client.post("/api/export/json/batch", json={
            "analysisIds": [record_id, str(self.other_record.analysis_id)],
        })
        self.assertEqual(zipped.status_code, 200, zipped.text)
        with zipfile.ZipFile(io.BytesIO(zipped.content)) as archive:
            self.assertEqual(archive.namelist(), ["sample.json"])
        request = {"selectedColumns": ["filename", "stemArea"], "unit": "um", "scale": 1}
        for endpoint, body in [
            ("/api/excel/" + record_id, request),
            ("/api/excel/batch", dict(request, analysisIds=[record_id])),
            ("/api/excel/summary", dict(request, taskType="stem")),
        ]:
            response = await self.client.post(endpoint, json=body)
            self.assertEqual(response.status_code, 200, response.text)
            # 新版前端用 responseType:'blob' 接，并靠 Content-Disposition 取中文文件名
            self.assertEqual(response.headers["content-type"], XLSX_MEDIA_TYPE)
            self.assertIn("attachment;", response.headers["content-disposition"])
            self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])
            workbook = load_workbook(io.BytesIO(response.content))
            self.assertEqual(workbook.active.cell(2, 2).value, 16)
            workbook.close()

    async def test_async_export_reports_progress_then_serves_file(self):
        # 新版前端调 /excel/summary 时只传 selectedColumns / unit / asyncMode，
        # 不带 taskType —— 类型由后端按记录推断，这里就走这条真实路径。
        request = {"selectedColumns": ["filename", "stemArea"], "unit": "um", "scale": 1}
        response = await self.client.post("/api/excel/summary", json=dict(
            request, asyncMode=True))
        self.assertEqual(response.status_code, 200, response.text)
        submitted = response.json()
        self.assertTrue(submitted["success"])
        task_id = submitted["taskId"]

        # 后台线程可能要跑一会儿：轮询到终态为止
        status = {}
        for _ in range(100):
            polled = await self.client.get("/api/excel/tasks/" + task_id)
            self.assertEqual(polled.status_code, 200, polled.text)
            status = polled.json()
            if status["status"] in ("completed", "failed"):
                break
            await asyncio.sleep(0.05)
        self.assertEqual(status["status"], "completed", status)
        self.assertEqual(status["total"], 1)
        self.assertTrue(status["downloadUrl"].endswith(task_id))

        downloaded = await self.client.get("/api/excel/download/" + task_id)
        self.assertEqual(downloaded.status_code, 200, downloaded.text)
        self.assertEqual(downloaded.headers["content-type"], XLSX_MEDIA_TYPE)
        workbook = load_workbook(io.BytesIO(downloaded.content))
        self.assertEqual(workbook.active.cell(2, 2).value, 16)
        workbook.close()

        missing = await self.client.get("/api/excel/tasks/does-not-exist")
        self.assertEqual(missing.status_code, 404)

    async def test_models_endpoint_and_model_name_upload(self):
        response = await self.client.get("/api/analysis/models")
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        by_key = {item["key"]: item for item in payload["models"]}
        self.assertIn("stem", by_key)
        self.assertTrue(by_key["stem"]["path"])
        self.assertEqual(by_key["stem"]["group"], "stem")
        # 默认模型是 stem（列表首项），下拉框展示的是中文显示名
        self.assertEqual(payload["default"], by_key["stem"]["name"])

        stem_only = await self.client.get("/api/analysis/models?group=stem")
        self.assertEqual([item["key"] for item in stem_only.json()["models"]], ["stem"])
        bad_group = await self.client.get("/api/analysis/models?group=nope")
        self.assertEqual(bad_group.status_code, 400)

        # 前端把 select 的 value（中文显示名）原样作为 model_name 回传
        display_name = by_key["stem"]["name"]
        buffer = io.BytesIO()
        Image.new("RGB", (16, 16), "white").save(buffer, format="PNG")
        upload = await self.client.post("/api/analysis/upload", files={
            "file": ("model.png", buffer.getvalue(), "image/png"),
        }, data={"model_name": display_name, "batch_id": "batch_1_abc", "batch_name": "我的批次"})
        self.assertEqual(upload.status_code, 200, upload.text)
        body = upload.json()
        self.assertEqual(body["task_type"], "stem")
        self.assertEqual(body["model_used"], display_name)
        record = self.session.get(Analysis, uuid.UUID(body["analysis_id"]))
        self.assertEqual(record.task_type, "stem")
        self.assertEqual(record.model_used, display_name)

        # 同一个 client batch_id 的第二张图应复用批次，历史里带上 batchName
        second = await self.client.post("/api/analysis/upload", files={
            "file": ("model2.png", buffer.getvalue(), "image/png"),
        }, data={"model_name": display_name, "batch_id": "batch_1_abc", "batch_name": "我的批次"})
        self.assertEqual(second.json()["batch_id"], body["batch_id"])

        history = await self.client.get("/api/analysis/history")
        entries = {item["analysisId"]: item for item in history.json()}
        entry = entries[body["analysis_id"]]
        self.assertEqual(entry["modelUsed"], display_name)
        self.assertIsNone(entry["errorMessage"])
        self.assertEqual(entry["batchName"], "我的批次")

    async def test_cross_user_access_is_rejected_and_delete_removes_owned_files(self):
        foreign_id = str(self.other_record.analysis_id)
        for endpoint in ["/api/export/json/" + foreign_id, "/api/export/json/preview/" + foreign_id]:
            response = await self.client.get(endpoint)
            self.assertEqual(response.status_code, 404, response.text)
        forbidden = await self.client.delete("/api/analysis/delete/" + foreign_id)
        self.assertEqual(forbidden.status_code, 404, forbidden.text)
        original = Path(self.owner_record.original_file_path)
        own_id = str(self.owner_record.analysis_id)
        response = await self.client.post("/api/analysis/delete/batch", json={"analysisIds": [own_id, foreign_id]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["deleted_count"], 1)
        self.assertFalse(original.exists())
        self.assertFalse(Path(preview_path_for(str(original))).exists())
        self.assertTrue(Path(self.other_record.original_file_path).exists())

    async def test_worker_persists_success_and_failure_timestamps(self):
        record = self.owner_record
        kwargs = dict(analysis_id=record.analysis_id, original_file_path=record.original_file_path,
                      user_id=str(record.user_id), original_filename="sample.png", task_type="stem")
        with patch.object(analysis_service, "run_system", return_value=(
                record.annotated_image_path, record.result_json_path)):
            analysis_service.run_full_analysis(**kwargs)
        self.session.refresh(record)
        self.assertEqual(record.status, "completed")
        self.assertIsNotNone(record.started_at)
        self.assertIsNotNone(record.finished_at)
        with patch.object(analysis_service, "run_system", side_effect=RuntimeError("test inference failure")):
            with self.assertRaisesRegex(RuntimeError, "test inference failure"):
                analysis_service.run_full_analysis(**kwargs)
        self.session.refresh(record)
        self.assertEqual(record.status, "failed")
        self.assertIsNotNone(record.finished_at)

    async def test_legacy_sync_export_remains_user_scoped(self):
        from fastapi import HTTPException
        result = excel.excel_service.load_analysis_data(
            str(self.owner_record.analysis_id), str(self.owner.id), scale=1,
        )
        self.assertEqual(result["stemArea"], 16)
        with self.assertRaises(HTTPException) as error:
            excel.excel_service.load_analysis_data(str(self.other_record.analysis_id), str(self.owner.id))
        self.assertEqual(error.exception.status_code, 404)
